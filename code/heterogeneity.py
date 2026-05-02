from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
import statsmodels.api as sm

try:
    from docx import Document
except Exception:
    Document = None


PROJECT_ROOT = Path(__file__).resolve().parents[1]
ROOT = PROJECT_ROOT / "aggregateData"
INPUT_FILE = ROOT / "TS_CSI_300_PCA.xlsx"
OUT_DATA_FILE = ROOT / "TS_CSI_300_PCA_heterogeneity.xlsx"
OUT_GRID_FILE = ROOT / "TS_CSI_300_PCA_heterogeneity_grid.xlsx"
OUT_WORD_FILE = ROOT / "Heterogeneity_Regression_Table.docx"

Y_COL = "Rexcess_lead1"
Y_ABS_COL = "Abs_Rexcess_lead1"
SENT_COL = "PCAMS"
CONTROL_COLS: list[str] = []
MONTH_COL = "Month"
Q_LOW_GRID = np.round(np.arange(0.05, 0.35 + 1e-9, 0.05), 2)
Q_HIGH_GRID = np.round(np.arange(0.95, 0.65 - 1e-9, -0.05), 2)
HAC_LAGS = 4
SIG_P_MAIN = 0.10
SIG_P_STRICT = 0.05


def normalize_month(series: pd.Series) -> pd.Series:
    if pd.api.types.is_datetime64_any_dtype(series):
        return series

    s = series.astype("string").str.strip()
    s = s.str.replace(r"\.0$", "", regex=True)
    s = s.str.replace("/", "-", regex=False)
    out = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")

    mask = s.str.fullmatch(r"\d{6}")
    if mask.any():
        out.loc[mask] = pd.to_datetime(s[mask], format="%Y%m", errors="coerce")

    mask = s.str.fullmatch(r"\d{4}-\d{1,2}")
    if mask.any():
        parts = s[mask].str.split("-", expand=True)
        tmp = parts[0] + "-" + parts[1].str.zfill(2)
        out.loc[mask] = pd.to_datetime(tmp, format="%Y-%m", errors="coerce")

    mask = s.str.fullmatch(r"\d{8}")
    if mask.any():
        out.loc[mask] = pd.to_datetime(s[mask], format="%Y%m%d", errors="coerce")

    mask = s.str.fullmatch(r"\d{4}-\d{2}-\d{2}")
    if mask.any():
        out.loc[mask] = pd.to_datetime(s[mask], format="%Y-%m-%d", errors="coerce")

    mask = out.isna() & s.notna()
    if mask.any():
        out.loc[mask] = pd.to_datetime(s[mask], errors="coerce")
    return out


def pval_stars(p: float) -> str:
    if p < 0.01:
        return "***"
    if p < 0.05:
        return "**"
    if p < 0.10:
        return "*"
    return ""


def as_series(index: pd.Index, values: np.ndarray) -> pd.Series:
    return pd.Series(np.asarray(values), index=index)


def write_word(
    out_file: Path,
    ols_res: sm.regression.linear_model.RegressionResultsWrapper,
    hac_res: sm.regression.linear_model.RegressionResultsWrapper,
    q_low: float,
    q_high: float,
    q_low_value: float,
    q_high_value: float,
    nobs: int,
) -> None:
    if Document is None:
        print("[Word] python-docx not installed. Skipping Word export.")
        return

    doc = Document()
    doc.add_heading("Heterogeneity Regression", level=1)
    model_rhs = f"{SENT_COL} + {SENT_COL}_High + {SENT_COL}_Low"
    if CONTROL_COLS:
        model_rhs += " + " + " + ".join(CONTROL_COLS)
    doc.add_paragraph(f"Model: {Y_ABS_COL} ~ {model_rhs}")
    doc.add_paragraph(
        f"High: {SENT_COL} >= {q_high:.0%} quantile ({q_high_value:.6f})"
    )
    doc.add_paragraph(
        f"Low: {SENT_COL} <= {q_low:.0%} quantile ({q_low_value:.6f})"
    )
    doc.add_paragraph(f"Observations used: {nobs}")
    doc.add_paragraph("Significance: *** p<0.01, ** p<0.05, * p<0.10")

    # HAC(Newey-West) table as the reported inference
    doc.add_heading(f"Newey-West (HAC, lags={HAC_LAGS})", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Variable"
    table.cell(0, 1).text = "Coef"
    table.cell(0, 2).text = "HAC Std.Err"
    table.cell(0, 3).text = "HAC t"
    table.cell(0, 4).text = "HAC p"

    hac_bse = as_series(ols_res.params.index, hac_res.bse)
    hac_t = as_series(ols_res.params.index, hac_res.tvalues)
    hac_p = as_series(ols_res.params.index, hac_res.pvalues)

    for v in ols_res.params.index:
        row = table.add_row().cells
        row[0].text = v
        row[1].text = f"{ols_res.params[v]:.4f}{pval_stars(float(hac_p[v]))}"
        row[2].text = f"{float(hac_bse[v]):.4f}"
        row[3].text = f"{float(hac_t[v]):.4f}"
        row[4].text = f"{float(hac_p[v]):.4f}"

    row = table.add_row().cells
    row[0].text = "N"
    row[1].text = str(int(ols_res.nobs))
    row[2].text = "R2"
    row[3].text = f"{ols_res.rsquared:.4f}"
    row[4].text = f"Adj.R2={ols_res.rsquared_adj:.4f}"

    try:
        doc.save(out_file)
        print("Saved Word table:", out_file)
    except PermissionError as exc:
        raise PermissionError(
            f"Cannot overwrite {out_file}. Close the file and rerun."
        ) from exc


def main() -> None:
    df = pd.read_excel(INPUT_FILE)
    required = [MONTH_COL, Y_COL, SENT_COL, *CONTROL_COLS]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise KeyError(f"Missing columns in {INPUT_FILE}: {missing}")

    # Keep time-series order
    df["_month"] = normalize_month(df[MONTH_COL])
    df = df.sort_values("_month").reset_index(drop=True)

    # Ensure numeric
    df[Y_COL] = pd.to_numeric(df[Y_COL], errors="coerce")
    df[Y_ABS_COL] = df[Y_COL].abs()
    df[SENT_COL] = pd.to_numeric(df[SENT_COL], errors="coerce")
    for c in CONTROL_COLS:
        df[c] = pd.to_numeric(df[c], errors="coerce")

    high_name = f"{SENT_COL}_High"
    low_name = f"{SENT_COL}_Low"
    regressors = [SENT_COL, high_name, low_name, *CONTROL_COLS]

    grid_rows: list[dict] = []
    model_cache: dict[tuple[float, float], dict] = {}

    for q_high in Q_HIGH_GRID:
        for q_low in Q_LOW_GRID:
            if q_low >= q_high:
                continue

            tmp = df.copy()
            q_low_value = float(tmp[SENT_COL].quantile(float(q_low)))
            q_high_value = float(tmp[SENT_COL].quantile(float(q_high)))
            tmp["High"] = np.where(tmp[SENT_COL] >= q_high_value, 1, 0)
            tmp["Low"] = np.where(tmp[SENT_COL] <= q_low_value, 1, 0)
            tmp[high_name] = tmp[SENT_COL] * tmp["High"]
            tmp[low_name] = tmp[SENT_COL] * tmp["Low"]

            reg_df = tmp[[Y_ABS_COL, *regressors]].dropna().copy()
            if reg_df.empty:
                continue

            y = reg_df[Y_ABS_COL]
            X = sm.add_constant(reg_df[regressors])
            ols_res = sm.OLS(y, X).fit()
            hac_res = ols_res.get_robustcov_results(cov_type="HAC", maxlags=HAC_LAGS)
            hac_p = as_series(ols_res.params.index, hac_res.pvalues)

            row = {
                "q_high": float(q_high),
                "q_low": float(q_low),
                "q_high_value": q_high_value,
                "q_low_value": q_low_value,
                "nobs": int(ols_res.nobs),
                "r2": float(ols_res.rsquared),
                "adj_r2": float(ols_res.rsquared_adj),
                "f_pvalue": float(ols_res.f_pvalue),
                "coef_PCAMS": float(ols_res.params.get(SENT_COL, np.nan)),
                "p_PCAMS": float(ols_res.pvalues.get(SENT_COL, np.nan)),
                "hac_p_PCAMS": float(hac_p.get(SENT_COL, np.nan)),
                "coef_PCAMS_High": float(ols_res.params.get(high_name, np.nan)),
                "p_PCAMS_High": float(ols_res.pvalues.get(high_name, np.nan)),
                "hac_p_PCAMS_High": float(hac_p.get(high_name, np.nan)),
                "coef_PCAMS_Low": float(ols_res.params.get(low_name, np.nan)),
                "p_PCAMS_Low": float(ols_res.pvalues.get(low_name, np.nan)),
                "hac_p_PCAMS_Low": float(hac_p.get(low_name, np.nan)),
            }
            pvals = [row["hac_p_PCAMS"], row["hac_p_PCAMS_High"], row["hac_p_PCAMS_Low"]]
            row["sig_count_main"] = int(sum(p <= SIG_P_MAIN for p in pvals))
            row["sig_count_strict"] = int(sum(p <= SIG_P_STRICT for p in pvals))
            row["all_sig_main"] = int(all(p <= SIG_P_MAIN for p in pvals))
            row["all_sig_strict"] = int(all(p <= SIG_P_STRICT for p in pvals))
            row["max_hac_p"] = float(max(pvals))
            row["sum_hac_p"] = float(sum(pvals))
            grid_rows.append(row)
            model_cache[(float(q_high), float(q_low))] = {
                "q_high_value": q_high_value,
                "q_low_value": q_low_value,
                "ols_res": ols_res,
                "hac_res": hac_res,
                "nobs": int(ols_res.nobs),
                "df": tmp.copy(),
            }

    if not grid_rows:
        raise ValueError("No valid regression results from quantile grid search.")

    grid_df = pd.DataFrame(grid_rows)
    candidates = grid_df[grid_df["all_sig_main"] == 1].copy()
    if candidates.empty:
        candidates = grid_df.copy()
    candidates = candidates.sort_values(
        [
            "all_sig_main",
            "sig_count_main",
            "all_sig_strict",
            "sig_count_strict",
            "max_hac_p",
            "adj_r2",
        ],
        ascending=[False, False, False, False, True, False],
    ).reset_index(drop=True)
    best_row = candidates.iloc[0]
    best_key = (float(best_row["q_high"]), float(best_row["q_low"]))
    best_model = model_cache[best_key]

    # Save full grid sorted by significance-first rule
    grid_df = grid_df.sort_values(
        [
            "all_sig_main",
            "sig_count_main",
            "all_sig_strict",
            "sig_count_strict",
            "max_hac_p",
            "adj_r2",
        ],
        ascending=[False, False, False, False, True, False],
    ).reset_index(drop=True)
    grid_df.to_excel(OUT_GRID_FILE, index=False)
    print("Saved grid search results:", OUT_GRID_FILE)

    print(
        f"\n=== Top 10 Grid Results (significance-first; main p<={SIG_P_MAIN:.2f}, strict p<={SIG_P_STRICT:.2f}) ==="
    )
    print(
        grid_df[
            [
                "q_high",
                "q_low",
                "all_sig_main",
                "sig_count_main",
                "all_sig_strict",
                "sig_count_strict",
                "max_hac_p",
                "adj_r2",
                "r2",
                "f_pvalue",
                "hac_p_PCAMS",
                "hac_p_PCAMS_High",
                "hac_p_PCAMS_Low",
            ]
        ]
        .head(10)
        .to_string(index=False)
    )

    print("\n=== Best Grid OLS Summary ===")
    print(best_model["ols_res"].summary())
    print(f"\n=== Best Grid HAC Coef Table (lags={HAC_LAGS}) ===")
    best_hac = best_model["hac_res"]
    print(
        pd.DataFrame(
            {
                "coef": best_model["ols_res"].params.values,
                "hac_se": best_hac.bse,
                "hac_t": best_hac.tvalues,
                "hac_p": best_hac.pvalues,
            },
            index=best_model["ols_res"].params.index,
        ).to_string()
    )

    # Save interaction DataFrame for best grid point
    out_cols = [
        MONTH_COL,
        SENT_COL,
        "High",
        "Low",
        high_name,
        low_name,
        *CONTROL_COLS,
        Y_COL,
        Y_ABS_COL,
    ]
    out_df = best_model["df"][out_cols].copy()
    try:
        out_df.to_excel(OUT_DATA_FILE, index=False)
        print("Saved interaction DataFrame (best grid point):", OUT_DATA_FILE)
    except PermissionError:
        alt_data = OUT_DATA_FILE.with_name(OUT_DATA_FILE.stem + "_new" + OUT_DATA_FILE.suffix)
        out_df.to_excel(alt_data, index=False)
        print("Data file is open/locked. Saved to:", alt_data)

    write_word(
        out_file=OUT_WORD_FILE,
        ols_res=best_model["ols_res"],
        hac_res=best_model["hac_res"],
        q_low=float(best_row["q_low"]),
        q_high=float(best_row["q_high"]),
        q_low_value=float(best_model["q_low_value"]),
        q_high_value=float(best_model["q_high_value"]),
        nobs=best_model["nobs"],
    )


if __name__ == "__main__":
    main()
