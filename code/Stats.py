from __future__ import annotations

from pathlib import Path
import warnings

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

try:
    from docx import Document
except Exception:
    Document = None


PROJECT_ROOT = Path(__file__).resolve().parents[1]
ROOT = PROJECT_ROOT / "aggregateData"
FILE_300 = ROOT / "TS_CSI_300_PCA.xlsx"
FILE_ALL = ROOT / "TS_CSI_ALL_PCA.xlsx"
OUT_FIG = ROOT / "PCAMS_vs_Rexcess_2x2.png"

MONTH_COL = "Month"
REX_COL = "Rexcess"
PCAMS_COL = "PCAMS"
EXPCAMS_COL = "exPCAMS"

# Ensure Chinese characters render
plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS"]
plt.rcParams["axes.unicode_minus"] = False


def normalize_month(series: pd.Series) -> pd.Series:
    if pd.api.types.is_datetime64_any_dtype(series):
        return series.dt.to_period("M")

    s = series.astype("string").str.strip()
    s = s.str.replace(r"\.0$", "", regex=True)
    s = (
        s.str.replace("/", "-", regex=False)
        .str.replace("年", "-", regex=False)
        .str.replace("月", "", regex=False)
    )

    out = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")

    mask = s.str.fullmatch(r"\d{6}")
    if mask.any():
        out.loc[mask] = pd.to_datetime(s[mask], format="%Y%m", errors="coerce")

    mask = s.str.fullmatch(r"\d{4}-\d{1,2}")
    if mask.any():
        parts = s[mask].str.split("-", expand=True)
        tmp = parts[0] + "-" + parts[1].str.zfill(2)
        out.loc[mask] = pd.to_datetime(tmp, format="%Y-%m", errors="coerce")

    mask = s.str.fullmatch(r"\d{8}")
    if mask.any():
        out.loc[mask] = pd.to_datetime(s[mask], format="%Y%m%d", errors="coerce")

    mask = s.str.fullmatch(r"\d{4}-\d{2}-\d{2}")
    if mask.any():
        out.loc[mask] = pd.to_datetime(s[mask], format="%Y-%m-%d", errors="coerce")

    mask = out.isna() & s.notna()
    if mask.any():
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            out.loc[mask] = pd.to_datetime(s[mask], errors="coerce")

    return out.dt.to_period("M")


def load_ts(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path)
    for col in [MONTH_COL, REX_COL, PCAMS_COL, EXPCAMS_COL]:
        if col not in df.columns:
            raise KeyError(f"Missing {col} in {path}")

    df[MONTH_COL] = normalize_month(df[MONTH_COL]).astype("string")
    df[REX_COL] = pd.to_numeric(df[REX_COL], errors="coerce")
    df[PCAMS_COL] = pd.to_numeric(df[PCAMS_COL], errors="coerce")
    df[EXPCAMS_COL] = pd.to_numeric(df[EXPCAMS_COL], errors="coerce")
    return df


def write_desc_table_word(df_300: pd.DataFrame, df_all: pd.DataFrame, out_path: Path) -> None:
    if Document is None:
        print("\n[Desc] python-docx not installed. Skipping Word export.")
        print("Install with: pip install python-docx")
        return

    # Variables: from CSI300 and only Rexcess from CSI All
    vars_300 = ["Rexcess", "NIA", "NIPO", "CEFP", "MMSent", "PB", "PE", "Turnover_diff", "PCAMS", "exPCAMS"]
    var_all = ["Rexcess"]

    # Validate columns
    missing_300 = [v for v in vars_300 if v not in df_300.columns]
    if missing_300:
        raise KeyError(f"Missing columns in TS_CSI_300_PCA.xlsx: {missing_300}")
    missing_all = [v for v in var_all if v not in df_all.columns]
    if missing_all:
        raise KeyError(f"Missing columns in TS_CSI_ALL_PCA.xlsx: {missing_all}")

    # Build descriptive stats
    def stats(series: pd.Series) -> dict:
        s = pd.to_numeric(series, errors="coerce").dropna()
        return {
            "N": int(s.count()),
            "Mean": s.mean(),
            "Std": s.std(ddof=1),
            "Min": s.min(),
            "P25": s.quantile(0.25),
            "Median": s.median(),
            "P75": s.quantile(0.75),
            "Max": s.max(),
        }

    rows = {}
    for v in vars_300:
        rows[f"CSI300_{v}"] = stats(df_300[v])
    rows["CSIA_Rexcess"] = stats(df_all["Rexcess"])

    doc = Document()
    doc.add_heading("描述性统计", level=1)
    doc.add_paragraph("样本：CSI300（多变量）与中证全指（仅Rexcess）")

    metrics = ["N", "Mean", "Std", "Min", "P25", "Median", "P75", "Max"]
    table = doc.add_table(rows=1 + len(rows), cols=1 + len(metrics))
    table.style = "Table Grid"

    # Header
    table.cell(0, 0).text = "变量"
    for j, m in enumerate(metrics, start=1):
        table.cell(0, j).text = m

    # Rows
    for i, (var, st) in enumerate(rows.items(), start=1):
        table.cell(i, 0).text = var
        for j, m in enumerate(metrics, start=1):
            val = st[m]
            if m == "N":
                table.cell(i, j).text = str(val)
            else:
                table.cell(i, j).text = f"{val:.4f}"

    doc.save(out_path)
    print("Saved descriptive table to:", out_path)


def zscore(series: pd.Series, mean: float, std: float) -> pd.Series:
    if std == 0 or np.isnan(std):
        return series * np.nan
    return (series - mean) / std


def add_corr_box(ax, corr: float) -> None:
    text = f"Corr(Sent_t, Rexcess_t+1) = {corr:.4f}"
    ax.text(
        0.02,
        0.98,
        text,
        transform=ax.transAxes,
        ha="left",
        va="top",
        fontsize=9,
        bbox=dict(boxstyle="round,pad=0.3", facecolor="white", alpha=0.7, edgecolor="none"),
    )


def plot_panel(
    ax,
    df: pd.DataFrame,
    sent_col: str,
    title: str,
    rex_lim: tuple[float, float],
    sent_lim: tuple[float, float],
    show_left_ylabel: bool,
    show_right_ylabel: bool,
) -> None:
    df = df[[MONTH_COL, REX_COL, sent_col]].dropna().copy()
    if df.empty:
        ax.set_xlabel(title, fontsize=11, labelpad=8)
        ax.text(0.5, 0.5, "No data", ha="center", va="center", transform=ax.transAxes)
        return

    df = df.sort_values(MONTH_COL)
    x = pd.PeriodIndex(df[MONTH_COL], freq="M").to_timestamp("M")
    rex = df[REX_COL]
    sent = df[sent_col]

    ax2 = ax.twinx()
    ax.plot(x, rex, color="#0B3D91", linewidth=1.2, alpha=0.9, label="Rexcess")
    ax2.plot(x, sent, color="#F28E2B", linewidth=1.6, alpha=0.95, label=sent_col)

    ax.set_xlabel(title, fontsize=11, labelpad=8)
    if show_left_ylabel:
        ax.set_ylabel("超额收益", color="#0B3D91")
    else:
        ax.set_ylabel("")
        ax.tick_params(axis="y", left=False, labelleft=False)
    if show_right_ylabel:
        ax2.set_ylabel("情绪因子(Z)", color="#F28E2B")
    else:
        ax2.set_ylabel("")
        ax2.tick_params(axis="y", right=False, labelright=False)

    ax.set_ylim(*rex_lim)
    ax2.set_ylim(*sent_lim)

    corr = pd.Series(sent).corr(rex.shift(-1))
    add_corr_box(ax, corr)

    ax.grid(True, axis="y", alpha=0.2)


def main() -> None:
    df_300 = load_ts(FILE_300)
    df_all = load_ts(FILE_ALL)

    # Descriptive statistics to Word
    out_desc = ROOT / "Descriptive_Stats.docx"
    write_desc_table_word(df_300, df_all, out_desc)

    # Global z-score for sentiment factors to keep consistent axis
    pcams_all = pd.concat([df_300[PCAMS_COL], df_all[PCAMS_COL]], ignore_index=True)
    expc_all = pd.concat([df_300[EXPCAMS_COL], df_all[EXPCAMS_COL]], ignore_index=True)
    pcams_mean, pcams_std = pcams_all.mean(), pcams_all.std(ddof=0)
    expc_mean, expc_std = expc_all.mean(), expc_all.std(ddof=0)

    df_300["PCAMS_z"] = zscore(df_300[PCAMS_COL], pcams_mean, pcams_std)
    df_all["PCAMS_z"] = zscore(df_all[PCAMS_COL], pcams_mean, pcams_std)
    df_300["exPCAMS_z"] = zscore(df_300[EXPCAMS_COL], expc_mean, expc_std)
    df_all["exPCAMS_z"] = zscore(df_all[EXPCAMS_COL], expc_mean, expc_std)

    # Global y-limits
    rex_all = pd.concat([df_300[REX_COL], df_all[REX_COL]], ignore_index=True).dropna()
    rex_min, rex_max = rex_all.min(), rex_all.max()
    rex_pad = (rex_max - rex_min) * 0.05 if rex_max != rex_min else 0.1
    rex_lim = (rex_min - rex_pad, rex_max + rex_pad)

    sent_all = pd.concat(
        [df_300["PCAMS_z"], df_all["PCAMS_z"], df_300["exPCAMS_z"], df_all["exPCAMS_z"]],
        ignore_index=True,
    ).dropna()
    sent_min, sent_max = sent_all.min(), sent_all.max()
    sent_pad = (sent_max - sent_min) * 0.05 if sent_max != sent_min else 0.5
    sent_lim = (sent_min - sent_pad, sent_max + sent_pad)

    fig, axes = plt.subplots(2, 2, figsize=(14, 9), sharex=False)

    plot_panel(
        axes[0, 0],
        df_300.rename(columns={"exPCAMS_z": "Sent_z"}),
        "Sent_z",
        "图A: 传统情绪因子 vs 沪深300超额收益",
        rex_lim,
        sent_lim,
        show_left_ylabel=True,
        show_right_ylabel=False,
    )

    plot_panel(
        axes[0, 1],
        df_all.rename(columns={"exPCAMS_z": "Sent_z"}),
        "Sent_z",
        "图B: 传统情绪因子 vs 中证全指超额收益",
        rex_lim,
        sent_lim,
        show_left_ylabel=False,
        show_right_ylabel=True,
    )

    plot_panel(
        axes[1, 0],
        df_300.rename(columns={"PCAMS_z": "Sent_z"}),
        "Sent_z",
        "图C: 综合市场情绪因子（含文本） vs 沪深300超额收益",
        rex_lim,
        sent_lim,
        show_left_ylabel=True,
        show_right_ylabel=False,
    )

    plot_panel(
        axes[1, 1],
        df_all.rename(columns={"PCAMS_z": "Sent_z"}),
        "Sent_z",
        "图D: 综合市场情绪因子（含文本） vs 中证全指超额收益",
        rex_lim,
        sent_lim,
        show_left_ylabel=False,
        show_right_ylabel=True,
    )

    fig.tight_layout()
    fig.subplots_adjust(hspace=0.45, wspace=0.15, bottom=0.12)
    fig.savefig(OUT_FIG, dpi=300)
    plt.show()
    print("Saved figure:", OUT_FIG)


if __name__ == "__main__":
    main()
