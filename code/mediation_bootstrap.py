from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
import statsmodels.api as sm
from scipy.stats import norm

try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except Exception:
    Document = None
    OxmlElement = None
    qn = None


ROOT = Path(__file__).resolve().parents[1] / "aggregateData"
DEFAULT_FILE = ROOT / "TS_CSI_300_PCA.xlsx"
DEFAULT_DOCX = ROOT / "Mediation_A1_PE_plus_Abs_to_fwd3mean_Table.docx"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Final dual mediation: A1_PE_plus_Abs_to_fwd3mean, with three-line table output."
    )
    parser.add_argument("--file", type=Path, default=DEFAULT_FILE, help="Input xlsx file.")
    parser.add_argument("--x", default="PCAMS", help="Independent variable.")
    parser.add_argument("--m1", default="PE", help="Mediator-1 variable (PE path).")
    parser.add_argument(
        "--m2",
        default="Abs_Rexcess",
        help="Mediator-2 variable (abs return path; built from y_base if missing).",
    )
    parser.add_argument("--y_base", default="Rexcess", help="Base return variable.")
    parser.add_argument("--month", default="Month", help="Time column for sorting.")
    parser.add_argument("--n_boot", type=int, default=5000, help="Bootstrap iterations.")
    parser.add_argument("--seed", type=int, default=2026, help="Random seed.")
    parser.add_argument("--out_docx", type=Path, default=DEFAULT_DOCX, help="Output Word table path.")
    parser.add_argument("--no_word", action="store_true", help="Skip Word export.")
    return parser.parse_args()


def normalize_month(series: pd.Series) -> pd.Series:
    if pd.api.types.is_datetime64_any_dtype(series):
        return series

    s = series.astype("string").str.strip()
    s = s.str.replace(r"\.0$", "", regex=True)
    s = s.str.replace("/", "-", regex=False)
    s = s.str.replace(r"[^0-9-]", "", regex=True)
    return pd.to_datetime(s, errors="coerce")


def fit_ols(y: pd.Series, x_df: pd.DataFrame) -> sm.regression.linear_model.RegressionResultsWrapper:
    x_design = sm.add_constant(x_df, has_constant="add")
    return sm.OLS(y, x_design).fit()


def future_mean(series: pd.Series, horizon: int) -> pd.Series:
    leads = [series.shift(-k) for k in range(1, horizon + 1)]
    return pd.concat(leads, axis=1).mean(axis=1)


def bootstrap_dual_indirect(
    df: pd.DataFrame,
    x_col: str,
    m1_col: str,
    m2_col: str,
    y_col: str,
    n_boot: int,
    seed: int,
) -> dict:
    rng = np.random.default_rng(seed)
    n = len(df)
    indirect1_vals = np.full(n_boot, np.nan, dtype=float)
    indirect2_vals = np.full(n_boot, np.nan, dtype=float)
    total_vals = np.full(n_boot, np.nan, dtype=float)

    for i in range(n_boot):
        idx = rng.integers(0, n, size=n)
        sample = df.iloc[idx]
        try:
            med1 = fit_ols(sample[m1_col], sample[[x_col]])
            med2 = fit_ols(sample[m2_col], sample[[x_col]])
            out = fit_ols(sample[y_col], sample[[x_col, m1_col, m2_col]])
            a1 = float(med1.params[x_col])
            b1 = float(out.params[m1_col])
            a2 = float(med2.params[x_col])
            b2 = float(out.params[m2_col])
            indirect1_vals[i] = a1 * b1
            indirect2_vals[i] = a2 * b2
            total_vals[i] = indirect1_vals[i] + indirect2_vals[i]
        except Exception:
            continue

    mask = np.isfinite(indirect1_vals) & np.isfinite(indirect2_vals) & np.isfinite(total_vals)
    i1 = indirect1_vals[mask]
    i2 = indirect2_vals[mask]
    it = total_vals[mask]
    if i1.size == 0:
        raise RuntimeError("Bootstrap failed: no valid resamples.")

    i1_ci = np.percentile(i1, [2.5, 97.5])
    i2_ci = np.percentile(i2, [2.5, 97.5])
    it_ci = np.percentile(it, [2.5, 97.5])

    return {
        "valid_bootstrap_samples": int(i1.size),
        "boot_m1_mean": float(np.mean(i1)),
        "boot_m1_ci_low": float(i1_ci[0]),
        "boot_m1_ci_high": float(i1_ci[1]),
        "boot_m1_significant": not (i1_ci[0] <= 0 <= i1_ci[1]),
        "boot_m2_mean": float(np.mean(i2)),
        "boot_m2_ci_low": float(i2_ci[0]),
        "boot_m2_ci_high": float(i2_ci[1]),
        "boot_m2_significant": not (i2_ci[0] <= 0 <= i2_ci[1]),
        "boot_total_mean": float(np.mean(it)),
        "boot_total_ci_low": float(it_ci[0]),
        "boot_total_ci_high": float(it_ci[1]),
        "boot_total_significant": not (it_ci[0] <= 0 <= it_ci[1]),
    }


def fmt_ci(low: float, high: float) -> str:
    return f"[{low:.4f}, {high:.4f}]"


def print_three_line_table(rows: list[dict]) -> None:
    print("\n=== Mediation Effects (Three-line Table) ===")
    top = "-" * 106
    print(top)
    print(f"{'Path':<30}{'Point(a*b)':>14}{'Bootstrap Mean':>18}{'95% CI':>28}{'Significant':>16}")
    print(top)
    for r in rows:
        print(
            f"{r['path']:<30}"
            f"{r['point']:>14.4f}"
            f"{r['boot_mean']:>18.4f}"
            f"{fmt_ci(r['ci_low'], r['ci_high']):>28}"
            f"{str(r['significant']):>16}"
        )
    print(top)


def _set_cell_border(cell, **kwargs) -> None:
    if OxmlElement is None or qn is None:
        return
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge, edge_data in kwargs.items():
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key in ["val", "sz", "space", "color"]:
            if key in edge_data:
                element.set(qn(f"w:{key}"), str(edge_data[key]))


def apply_three_line_style(table) -> None:
    rows = len(table.rows)
    cols = len(table.columns)

    # Clear all borders first
    for i in range(rows):
        for j in range(cols):
            _set_cell_border(
                table.cell(i, j),
                left={"val": "nil"},
                right={"val": "nil"},
                top={"val": "nil"},
                bottom={"val": "nil"},
            )

    # Top line and header bottom line
    for j in range(cols):
        _set_cell_border(
            table.cell(0, j),
            top={"val": "single", "sz": 12, "color": "000000"},
            bottom={"val": "single", "sz": 12, "color": "000000"},
        )

    # Bottom line
    for j in range(cols):
        _set_cell_border(
            table.cell(rows - 1, j),
            bottom={"val": "single", "sz": 12, "color": "000000"},
        )


def write_three_line_word(
    out_path: Path,
    model_info: dict,
    final_rows: list[dict],
    effect_rows: list[dict],
) -> None:
    if Document is None:
        print("[Word] python-docx not installed. Skip Word export.")
        print("Install with: pip install python-docx")
        return

    doc = Document()
    doc.add_heading("Dual Mediation Result", level=1)
    doc.add_paragraph("Table 1. Final Model Summary and Path Coefficients")
    table1 = doc.add_table(rows=1 + len(final_rows), cols=2)
    table1.cell(0, 0).text = "Item"
    table1.cell(0, 1).text = "Value"
    for i, r in enumerate(final_rows, start=1):
        table1.cell(i, 0).text = r["item"]
        table1.cell(i, 1).text = r["value"]
    apply_three_line_style(table1)

    doc.add_paragraph("")
    doc.add_paragraph("Table 2. Bootstrap Mediation Effects")
    table2 = doc.add_table(rows=1 + len(effect_rows), cols=5)
    table2.cell(0, 0).text = "Path"
    table2.cell(0, 1).text = "Point (a*b)"
    table2.cell(0, 2).text = "Bootstrap Mean"
    table2.cell(0, 3).text = "95% CI"
    table2.cell(0, 4).text = "Significant"
    for i, r in enumerate(effect_rows, start=1):
        table2.cell(i, 0).text = r["path"]
        table2.cell(i, 1).text = f"{r['point']:.4f}"
        table2.cell(i, 2).text = f"{r['boot_mean']:.4f}"
        table2.cell(i, 3).text = fmt_ci(r["ci_low"], r["ci_high"])
        table2.cell(i, 4).text = "Yes" if r["significant"] else "No"
    apply_three_line_style(table2)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print("Saved two three-line Word tables to:", out_path)


def main() -> None:
    args = parse_args()
    if not args.file.exists():
        raise FileNotFoundError(f"Input file not found: {args.file}")

    raw = pd.read_excel(args.file)
    required = [args.month, args.x, args.m1, args.y_base]
    for col in required:
        if col not in raw.columns:
            raise KeyError(f"Missing required column: {col}")

    df = raw.copy()
    df["_month_parsed"] = normalize_month(df[args.month])
    df = df.sort_values("_month_parsed").reset_index(drop=True)
    df[args.x] = pd.to_numeric(df[args.x], errors="coerce")
    df[args.m1] = pd.to_numeric(df[args.m1], errors="coerce")
    df[args.y_base] = pd.to_numeric(df[args.y_base], errors="coerce")

    if args.m2 not in df.columns:
        if args.m2 == "Abs_Rexcess":
            df[args.m2] = df[args.y_base].abs()
        else:
            raise KeyError(f"Missing mediator-2 column: {args.m2}")
    else:
        df[args.m2] = pd.to_numeric(df[args.m2], errors="coerce")

    y_col = "Rexcess_fwd3_mean"
    df[y_col] = future_mean(df[args.y_base], horizon=3)

    model_df = df[[args.x, args.m1, args.m2, y_col]].dropna().reset_index(drop=True)
    n_raw = len(df)
    n_used = len(model_df)
    if n_used < 20:
        raise ValueError(f"Too few valid observations after dropna: {n_used}")

    med1 = fit_ols(model_df[args.m1], model_df[[args.x]])
    med2 = fit_ols(model_df[args.m2], model_df[[args.x]])
    out = fit_ols(model_df[y_col], model_df[[args.x, args.m1, args.m2]])
    total = fit_ols(model_df[y_col], model_df[[args.x]])

    a1 = float(med1.params[args.x])
    b1 = float(out.params[args.m1])
    a2 = float(med2.params[args.x])
    b2 = float(out.params[args.m2])
    c_prime = float(out.params[args.x])
    c_total = float(total.params[args.x])

    indirect_m1 = a1 * b1
    indirect_m2 = a2 * b2
    indirect_total = indirect_m1 + indirect_m2

    se_a1 = float(med1.bse[args.x])
    se_b1 = float(out.bse[args.m1])
    se_a2 = float(med2.bse[args.x])
    se_b2 = float(out.bse[args.m2])
    sobel1_se = np.sqrt((b1 ** 2) * (se_a1 ** 2) + (a1 ** 2) * (se_b1 ** 2))
    sobel2_se = np.sqrt((b2 ** 2) * (se_a2 ** 2) + (a2 ** 2) * (se_b2 ** 2))
    sobel1_z = np.nan if sobel1_se == 0 or np.isnan(sobel1_se) else indirect_m1 / sobel1_se
    sobel2_z = np.nan if sobel2_se == 0 or np.isnan(sobel2_se) else indirect_m2 / sobel2_se
    sobel1_p = np.nan if np.isnan(sobel1_z) else 2 * (1 - norm.cdf(abs(sobel1_z)))
    sobel2_p = np.nan if np.isnan(sobel2_z) else 2 * (1 - norm.cdf(abs(sobel2_z)))

    boot = bootstrap_dual_indirect(
        df=model_df,
        x_col=args.x,
        m1_col=args.m1,
        m2_col=args.m2,
        y_col=y_col,
        n_boot=args.n_boot,
        seed=args.seed,
    )

    print("\n=== Final Model: A1_PE_plus_Abs_to_fwd3mean ===")
    print(f"Input file: {args.file}")
    print(f"Rows before dropna: {n_raw}")
    print(f"Rows used: {n_used}")
    print(f"X: {args.x}, M1: {args.m1}, M2: {args.m2}, Y: {y_col}")
    print(
        "Path coefficients: "
        f"a1={a1:.4f}(p={float(med1.pvalues[args.x]):.4f}), "
        f"b1={b1:.4f}(p={float(out.pvalues[args.m1]):.4f}), "
        f"a2={a2:.4f}(p={float(med2.pvalues[args.x]):.4f}), "
        f"b2={b2:.4f}(p={float(out.pvalues[args.m2]):.4f}), "
        f"c'={c_prime:.4f}(p={float(out.pvalues[args.x]):.4f}), "
        f"c={c_total:.4f}(p={float(total.pvalues[args.x]):.4f})"
    )
    print(f"Sobel: M1 z={sobel1_z:.4f}, p={sobel1_p:.4f}; M2 z={sobel2_z:.4f}, p={sobel2_p:.4f}")

    effect_rows = [
        {
            "path": "M1 path (PE: a1*b1)",
            "point": indirect_m1,
            "boot_mean": boot["boot_m1_mean"],
            "ci_low": boot["boot_m1_ci_low"],
            "ci_high": boot["boot_m1_ci_high"],
            "significant": boot["boot_m1_significant"],
        },
        {
            "path": "M2 path (|Rexcess|: a2*b2)",
            "point": indirect_m2,
            "boot_mean": boot["boot_m2_mean"],
            "ci_low": boot["boot_m2_ci_low"],
            "ci_high": boot["boot_m2_ci_high"],
            "significant": boot["boot_m2_significant"],
        },
        {
            "path": "Total indirect",
            "point": indirect_total,
            "boot_mean": boot["boot_total_mean"],
            "ci_low": boot["boot_total_ci_low"],
            "ci_high": boot["boot_total_ci_high"],
            "significant": boot["boot_total_significant"],
        },
    ]

    print_three_line_table(effect_rows)

    if not args.no_word:
        final_rows = [
            {
                "item": "a1 (X -> M1)",
                "value": f"{a1:.4f} (p={float(med1.pvalues[args.x]):.4f})",
            },
            {
                "item": "b1 (M1 -> Y | X, M2)",
                "value": f"{b1:.4f} (p={float(out.pvalues[args.m1]):.4f})",
            },
            {
                "item": "a2 (X -> M2)",
                "value": f"{a2:.4f} (p={float(med2.pvalues[args.x]):.4f})",
            },
            {
                "item": "b2 (M2 -> Y | X, M1)",
                "value": f"{b2:.4f} (p={float(out.pvalues[args.m2]):.4f})",
            },
            {
                "item": "c' (Direct)",
                "value": f"{c_prime:.4f} (p={float(out.pvalues[args.x]):.4f})",
            },
            {
                "item": "c (Total)",
                "value": f"{c_total:.4f} (p={float(total.pvalues[args.x]):.4f})",
            },
            {
                "item": "Sobel M1",
                "value": f"z={sobel1_z:.4f}, p={sobel1_p:.4f}",
            },
            {
                "item": "Sobel M2",
                "value": f"z={sobel2_z:.4f}, p={sobel2_p:.4f}",
            },
        ]
        model_info = {
            "x": args.x,
            "m1": args.m1,
            "m2": args.m2,
            "y": y_col,
            "n_raw": n_raw,
            "n_used": n_used,
            "n_boot": args.n_boot,
            "seed": args.seed,
        }
        write_three_line_word(args.out_docx, model_info, final_rows, effect_rows)


if __name__ == "__main__":
    main()
